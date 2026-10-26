from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

LABELS = {'match': 'Con logotipo', 'no_match': 'Sin coincidencia', 'review': 'Revisar',
          'unclassified': 'Sin clasificar', 'processed': 'Preparado', 'skipped': 'Omitido', 'pending': 'Pendiente'}


def write_report(job: dict, directory: Path) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = 'Resumen'
    summary.append(['Concepto', 'Valor'])
    for row in [('Trabajo', job['name']), ('Creado', job['created']),
                ('Archivos originales', sum(f['original'] for f in job['files'])),
                ('Archivos generados', sum(not f['original'] for f in job['files'])),
                ('Incidencias', len(job['errors'])),
                ('Alcance', 'Reconocimiento de logotipos, sin validación de firmas.'),
                ('Carpetas vacías', 'El navegador no transmite carpetas vacías.')]:
        summary.append(row)
    folders = workbook.create_sheet('Carpetas')
    folders.append(['Carpeta', 'Archivos originales', 'Archivos generados', 'Resultado'])
    for group in job['groups']:
        files = [f for f in job['files'] if f['group'] == group['name']]
        folders.append([group['name'], sum(f['original'] for f in files),
                        sum(not f['original'] for f in files), LABELS[group['state']]])
    files = workbook.create_sheet('Archivos')
    files.append(['Carpeta', 'Archivo', 'Formato', 'Origen', 'Resultado', 'Detalle', 'Páginas con coincidencia'])
    for file in job['files']:
        files.append([file['group'], file['relative'], file['kind'].upper(),
                      'Original' if file['original'] else 'Generado', LABELS.get(file['state'], file['state']),
                      file['message'], ', '.join(str(m['page']) for m in file.get('matches', []))])
    errors = workbook.create_sheet('Incidencias')
    errors.append(['Archivo', 'Detalle'])
    for error in job['errors']:
        errors.append([error['file'], error['message']])
    for sheet in workbook:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    # Explicit string type prevents formula interpretation of untrusted names.
                    cell.data_type = 's'
                    cell.value = ''.join(c for c in cell.value if ord(c) >= 32 or c in '\n\t')[:32000]
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='126F65')
            cell.font = Font(color='FFFFFF', bold=True)
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(65, max(18, max(len(str(c.value or '')) for c in column) + 2))
    path = directory / 'output' / 'reporte.xlsx'
    workbook.save(path)
    workbook.close()
    return path
